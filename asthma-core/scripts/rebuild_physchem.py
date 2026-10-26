# -*- coding: utf-8 -*-
"""
重建化合物理化性质表（修复 957 行整体右移一列的问题）

背景：
  数据库溯源/主要展示数据均为cctcm2.0/化合物理化性质和生物利用度属性-赵玉男.xlsx
  共 2,499 行数据（第 3 行起），其中 957 行（38%）整行数据右移一列：

      错位列            实际内容        证据
      LogS      ←  分子式            RDKit CalcMolFormula 一致率 100%
      LogD      ←  分子量            RDKit MolWt 一致率 85%
      LogP      ←  重原子数          RDKit GetNumHeavyAtoms 一致率 100%
      Pgp-inhibitor ← 形式电荷       —
      Pgp-substrate ← 立体中心数     —
      F(20%)    ←  杂原子数          一致率 88%
      Caco-2    ←  可旋转键          —
      MDCK      ←  TPSA              一致率 96%
      后 9 列   ←  全空

  原始 ADME 值（真实 LogS/LogD/Pgp/F20%/Caco-2/MDCK）已不可恢复。

策略：
  A. 可从结构精确计算的 12 个描述符 —— 全部用 RDKit 重算
  B. ADME 类 7 列 —— 错位行置空，交给下游模型中位数填补
  C. 保留原文件双行表头与 29 列结构（用 openpyxl 原地改，不用 to_excel）

用法：
  python -X utf8 scripts/rebuild_physchem.py            # dry-run
  python -X utf8 scripts/rebuild_physchem.py --apply    # 写回
"""
import argparse
import sys
from pathlib import Path

from rdkit import Chem, RDLogger
from rdkit.Chem import Crippen, rdMolDescriptors

sys.stdout.reconfigure(encoding="utf-8")
RDLogger.DisableLog("rdApp.*")

REPO = Path(__file__).resolve().parent.parent.parent.parent
SRC = REPO / "数据库溯源" / "主要展示数据均为cctcm2.0" / "化合物理化性质和生物利用度属性-赵玉男.xlsx"

# 数据从第 3 行开始（第 1 行分组标题、第 2 行列名）
DATA_START_ROW = 3
COL_NAME = 1     # Compound
COL_SMILES = 2   # smiles

# 列名 → 列索引（1-based，按原文件顺序）
COL = {
    "LogS": 3, "LogD": 4, "LogP": 5,
    "Pgp-inhibitor": 6, "Pgp-substrate": 7, "F(20%)": 8,
    "Caco-2 Permeability": 9, "MDCK Permeability (cm/s)": 10,
    "Num. H-bond acceptors": 11, "Num. H-bond donors": 12, "TPSA": 13,
    "Num. Rotatable bonds": 14, "Num. Rings": 15, "MaxRing": 16,
    "nHet": 17, "fChar": 18, "nRig": 19, "Flex": 20, "nStereo": 21,
}

RDKIT_COLS = ["LogP", "TPSA", "Num. H-bond acceptors", "Num. H-bond donors",
              "Num. Rotatable bonds", "Num. Rings", "MaxRing", "nHet",
              "fChar", "nRig", "Flex", "nStereo"]
ADME_COLS = ["LogS", "LogD", "Pgp-inhibitor", "Pgp-substrate",
             "F(20%)", "Caco-2 Permeability", "MDCK Permeability (cm/s)"]


def rdkit_descriptors(mol):
    rings = mol.GetRingInfo().AtomRings()
    n_bonds = mol.GetNumBonds()
    n_rot = rdMolDescriptors.CalcNumRotatableBonds(mol)
    return {
        "LogP": round(Crippen.MolLogP(mol), 3),
        "TPSA": round(rdMolDescriptors.CalcTPSA(mol), 2),
        "Num. H-bond acceptors": rdMolDescriptors.CalcNumHBA(mol),
        "Num. H-bond donors": rdMolDescriptors.CalcNumHBD(mol),
        "Num. Rotatable bonds": n_rot,
        "Num. Rings": rdMolDescriptors.CalcNumRings(mol),
        "MaxRing": max((len(r) for r in rings), default=0),
        "nHet": rdMolDescriptors.CalcNumHeteroatoms(mol),
        "fChar": Chem.GetFormalCharge(mol),
        "nRig": n_bonds - n_rot,
        "Flex": round((n_rot / n_bonds) if n_bonds > 0 else 0.0, 3),
        "nStereo": rdMolDescriptors.CalcNumAtomStereoCenters(mol),
    }


def is_formula_string(v):
    """判断单元格是否为分子式字符串（如 C17H24O11 / FeS2 / K）"""
    if v is None:
        return False
    s = str(v).strip()
    if not s or s in ("NA", "nan", "None"):
        return False
    if not s[:1].isupper():
        return False
    try:
        float(s)
        return False
    except ValueError:
        return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    import openpyxl

    print(f"源文件: {SRC}")
    wb = openpyxl.load_workbook(SRC)
    ws = wb[wb.sheetnames[0]]
    print(f"工作表: {ws.title}  rows={ws.max_row}  cols={ws.max_column}")

    shifted_rows = []
    parse_fail = 0
    fixed_cells = 0

    for r in range(DATA_START_ROW, ws.max_row + 1):
        smiles = ws.cell(r, COL_SMILES).value
        logs = ws.cell(r, COL["LogS"]).value
        if not smiles:
            parse_fail += 1
            continue

        shifted = is_formula_string(logs)
        mol = Chem.MolFromSmiles(str(smiles).strip())
        if mol is None:
            parse_fail += 1
            continue
        desc = rdkit_descriptors(mol)

        if shifted:
            shifted_rows.append(r)
            # 12 个结构描述符整列覆盖
            for c in RDKIT_COLS:
                ws.cell(r, COL[c]).value = desc[c]
                fixed_cells += 1
            # 7 个 ADME 列置空（原值不可信）
            for c in ADME_COLS:
                ws.cell(r, COL[c]).value = None
                fixed_cells += 1
        else:
            # 正常行：只补空缺
            for c in RDKIT_COLS:
                if ws.cell(r, COL[c]).value is None:
                    ws.cell(r, COL[c]).value = desc[c]
                    fixed_cells += 1

    print(f"错位行: {len(shifted_rows)}")
    print(f"SMILES 解析失败: {parse_fail}")
    print(f"待修改单元格: {fixed_cells}")

    if not args.apply:
        print("\n[dry-run] 未写文件")
        wb.close()
        return

    # 在第 2 行末尾加一列说明（不破坏原有 29 列结构）
    note_col = ws.max_column + 1
    ws.cell(2, note_col).value = "_data_quality_note"
    ws.cell(2, note_col).comment = None
    ws.cell(1, note_col).value = None
    # 逐行标注
    shifted_set = set(shifted_rows)
    for r in range(DATA_START_ROW, ws.max_row + 1):
        if r in shifted_set:
            ws.cell(r, note_col).value = "reconstructed_by_rdkit"
        elif ws.cell(r, COL_SMILES).value:
            ws.cell(r, note_col).value = "original"

    wb.save(SRC)
    wb.close()
    print(f"✓ 已写回 {SRC.name}（保留原双行表头与 29 列，新增第 {note_col} 列标注）")


if __name__ == "__main__":
    main()
